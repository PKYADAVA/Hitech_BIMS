"""Inventory > Items: the list page, bulk status, Last Updated, and importing
items from a sheet.

An import is matched to the masters by name and saved all or nothing, so a
half-imported sheet never has to be untangled by hand, and a sheet uploaded
twice does not make every item twice.
"""
import io
import json
from decimal import Decimal

import openpyxl
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse

from inventory.models import Item, ItemCategory, UnitOfMeasurement, Warehouse
from inventory.services.item_import import (ItemImportError, import_items,
                                            item_template_workbook,
                                            parse_item_upload)


class ItemsPageBase(TestCase):

    def setUp(self):
        self.feed = ItemCategory.objects.create(name="Broiler Feed")
        self.chicks = ItemCategory.objects.create(name="Day Old Chicks")
        self.bag = UnitOfMeasurement.objects.create(name="Bag", symbol="Bag")
        self.kilo = UnitOfMeasurement.objects.create(name="Kilogram", symbol="Kg")
        self.north = Warehouse.objects.create(name="Bahraich Warehouse")
        self.south = Warehouse.objects.create(name="Akbarpur Warehouse")
        self.starter = Item.objects.create(
            description="Starter Feed", category=self.feed, valuation_method="Weighted Average",
            standard_cost_per_unit=Decimal("42"), usage="Produced", source="Purchased",
            type="Raw Material", item_account="Expense", storage_uom=self.bag)
        self.user = get_user_model().objects.create_superuser("keeper", "k@x.com", "Str0ngPass!")
        self.client.force_login(self.user)

    def sheet(self, rows, headers=None, name="items.xlsx"):
        book = openpyxl.Workbook()
        book.active.title = "Items"
        book.active.append(headers or ["Description", "Category", "Valuation Method",
                                       "Standard Cost/Unit", "Usage", "Storage UOM",
                                       "Consumption UOM", "Kg per Bag", "HSN Code", "Warehouses"])
        for row in rows:
            book.active.append(list(row))
        buffer = io.BytesIO()
        book.save(buffer)
        return SimpleUploadedFile(name, buffer.getvalue())


class ListTests(ItemsPageBase):

    def test_the_page_renders_the_new_layout(self):
        response = self.client.get(reverse("items"))
        self.assertEqual(response.status_code, 200)
        for text in ('id="item-table"', 'id="itl-import-btn"', 'id="itl-f-warehouse"',
                     'id="itl-f-valuation"', 'id="itemFormModal"', "Last Updated"):
            self.assertContains(response, text)

    def test_the_list_carries_warehouse_names_and_last_updated(self):
        self.starter.warehouse.set([self.north, self.south])
        row = next(r for r in self.client.get(reverse("item_list")).json() if r["id"] == self.starter.id)
        self.assertEqual(sorted(row["warehouse_names"]), ["Akbarpur Warehouse", "Bahraich Warehouse"])
        self.assertTrue(row["updated_at"])

    def test_the_view_dialog_gets_readable_labels(self):
        detail = self.client.get(f"/item/{self.starter.id}/").json()
        self.assertEqual((detail["category_name"], detail["storage_uom_label"]), ("Broiler Feed", "Bag"))

    def test_saving_moves_last_updated(self):
        first = Item.objects.get(id=self.starter.id).updated_at
        self.client.post(reverse("items_toggle_active", args=[self.starter.id]))
        self.assertGreater(Item.objects.get(id=self.starter.id).updated_at, first)


class BulkStatusTests(ItemsPageBase):

    def test_ticked_items_are_made_inactive_together(self):
        grower = Item.objects.create(
            description="Grower Feed", category=self.feed, valuation_method="FIFO",
            standard_cost_per_unit=Decimal("40"), usage="Produced", source="Purchased",
            type="Raw Material", item_account="Expense")
        response = self.client.post(reverse("items_bulk_status_update"), json.dumps(
            {"ids": [self.starter.id, grower.id], "active": False}), content_type="application/json")
        self.assertEqual(response.json()["changed"], 2)
        self.assertFalse(Item.objects.filter(is_active=True).exists())

    def test_nothing_ticked_is_refused(self):
        response = self.client.post(reverse("items_bulk_status_update"), json.dumps(
            {"ids": [], "active": False}), content_type="application/json")
        self.assertEqual(response.status_code, 400)


class ImportPreviewTests(ItemsPageBase):

    def test_a_good_row_is_matched_to_the_masters(self):
        res = parse_item_upload(self.sheet([
            ("Grower Feed", "broiler feed", "Weighted Avg.", "40.50", "Production", "Bag", "kg",
             50, "23099010", "Bahraich Warehouse, Akbarpur Warehouse"),
        ]))
        row = res["rows"][0]
        self.assertEqual(row["action"], "create", row["message"])
        self.assertEqual((row["category"], row["valuation_method"], row["usage"]),
                         (self.feed.id, "Weighted Average", "Produced"))
        self.assertEqual((row["storage_uom"], row["consumption_uom"]), (self.bag.id, self.kilo.id))
        self.assertEqual(sorted(row["warehouses"]), sorted([self.north.id, self.south.id]))
        self.assertEqual((row["standard_cost_per_unit"], row["kg_per_bag"]), ("40.50", "50"))

    def test_all_means_every_warehouse(self):
        row = parse_item_upload(self.sheet([
            ("Grower Feed", "Broiler Feed", "FIFO", 40, "Sales", "", "", "", "", "All"),
        ]))["rows"][0]
        self.assertEqual(sorted(row["warehouses"]), sorted([self.north.id, self.south.id]))

    def test_each_problem_is_named_on_its_row(self):
        res = parse_item_upload(self.sheet([
            ("", "Broiler Feed", "FIFO", 40, "Sales", "", "", "", "", ""),
            ("Grower Feed", "Pet Food", "FIFO", 40, "Sales", "", "", "", "", ""),
            ("Grower Feed", "Broiler Feed", "Average", 40, "Sales", "", "", "", "", ""),
            ("Grower Feed", "Broiler Feed", "FIFO", "abc", "Sales", "", "", "", "", ""),
            ("Grower Feed", "Broiler Feed", "FIFO", 40, "Gift", "Tonne", "", "", "", "Delhi Store"),
            ("Starter Feed", "Broiler Feed", "FIFO", 40, "Sales", "", "", "", "", ""),
        ]))
        messages = [r["message"] for r in res["rows"]]
        self.assertIn("Description is missing", messages[0])
        self.assertIn("No category named Pet Food", messages[1])
        self.assertIn("Valuation method not recognised", messages[2])
        self.assertIn("not a number", messages[3])
        self.assertIn("Usage must be Produced or Sales", messages[4])
        self.assertIn("No unit named Tonne", messages[4])
        self.assertIn("No warehouse named Delhi Store", messages[4])
        self.assertIn(f"Already exists as {self.starter.item_code}", messages[5])
        self.assertEqual(res["counts"], {"create": 0, "error": 6})

    def test_the_same_item_twice_in_one_file_is_caught(self):
        res = parse_item_upload(self.sheet([
            ("Grower Feed", "Broiler Feed", "FIFO", 40, "Sales", "", "", "", "", ""),
            ("grower  feed", "Broiler Feed", "FIFO", 41, "Sales", "", "", "", "", ""),
        ]))
        self.assertEqual([r["action"] for r in res["rows"]], ["create", "error"])

    def test_a_csv_is_read_too(self):
        text = "Description,Category,Valuation Method,Standard Cost/Unit,Usage\nGrower Feed,Broiler Feed,LIFO,40,Sales\n"
        row = parse_item_upload(SimpleUploadedFile("items.csv", text.encode()))["rows"][0]
        self.assertEqual(row["action"], "create")

    def test_a_sheet_missing_required_columns_is_refused(self):
        with self.assertRaises(ItemImportError) as refused:
            parse_item_upload(self.sheet([("Grower Feed",)], headers=["Description"]))
        self.assertIn("Category", str(refused.exception))

    def test_the_page_previews_without_saving(self):
        response = self.client.post(reverse("items_import_data"), {"file": self.sheet([
            ("Grower Feed", "Broiler Feed", "FIFO", 40, "Sales", "", "", "", "", ""),
        ])})
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()["counts"]["create"], 1)
        self.assertFalse(Item.objects.filter(description="Grower Feed").exists())


class ImportApplyTests(ItemsPageBase):

    def good_rows(self):
        return parse_item_upload(self.sheet([
            ("Grower Feed", "Broiler Feed", "FIFO", 40, "Produced", "Bag", "Kg", 50, "23099010", "All"),
            ("Broiler Chicks", "Day Old Chicks", "Weighted Average", 35, "Sales", "", "", "", "", "Bahraich Warehouse"),
        ]))["rows"]

    def test_importing_creates_the_items_with_codes_and_warehouses(self):
        response = self.client.post(reverse("items_import_create"), json.dumps({"rows": self.good_rows()}),
                                    content_type="application/json")
        self.assertEqual(response.status_code, 201, response.content)
        self.assertEqual(response.json()["created"], 2)
        grower = Item.objects.get(description="Grower Feed")
        self.assertTrue(grower.item_code)
        self.assertEqual((grower.storage_uom, grower.kg_per_bag, grower.hsn_code),
                         (self.bag, Decimal("50.00"), "23099010"))
        self.assertEqual(grower.warehouse.count(), 2)
        self.assertEqual(list(Item.objects.get(description="Broiler Chicks").warehouse.all()), [self.north])

    def test_it_is_all_or_nothing(self):
        rows = self.good_rows()
        rows[1]["category"] = 999999
        with self.assertRaises(ItemImportError):
            import_items(rows)
        self.assertFalse(Item.objects.filter(description="Grower Feed").exists())

    def test_importing_the_same_rows_twice_is_refused(self):
        rows = self.good_rows()
        import_items(rows)
        with self.assertRaises(ItemImportError) as refused:
            import_items(rows)
        self.assertIn("already exists", str(refused.exception))
        self.assertEqual(Item.objects.filter(description="Grower Feed").count(), 1)

    def test_the_template_lists_what_the_columns_accept(self):
        response = self.client.get(reverse("items_template_data"))
        self.assertEqual(response.status_code, 200)
        book = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertEqual([c.value for c in book["Items"][1]][:5],
                         ["Description", "Category", "Valuation Method", "Standard Cost/Unit", "Usage"])
        lists = list(book["Lists"].iter_rows(values_only=True))
        self.assertIn("Broiler Feed", [r[0] for r in lists])
        self.assertIn("All", [r[2] for r in lists])
        self.assertIsNotNone(item_template_workbook())


class AccessMappingTests(TestCase):

    def test_the_new_endpoints_are_governed_by_the_items_tab(self):
        from user.access import derive_tab, resolve_action

        expected = {
            "items_import_data": "view",
            "items_template_data": "view",
            "items_import_create": "add",
            "items_bulk_status_update": "edit",
        }
        for name, action in expected.items():
            self.assertEqual(resolve_action(name) or derive_tab(name), ("items", action), name)
