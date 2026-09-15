"""Item import from Excel or CSV, for Inventory > Items.

A sheet of new items is read row by row and matched to the masters by name:
category, storage and consumption unit (name or symbol), and warehouses
(comma-separated, or All). Nothing is saved until the preview has been seen,
and then every row is saved or none, so a half-imported sheet never has to be
untangled by hand.
"""
import csv
import io
from decimal import Decimal, InvalidOperation

from django.db import transaction

from inventory.models import Item, ItemCategory, UnitOfMeasurement, Warehouse

MAX_ROWS = 2000

COLUMNS = ["Description", "Category", "Valuation Method", "Standard Cost/Unit", "Usage",
           "Storage UOM", "Consumption UOM", "Kg per Bag", "HSN Code", "Warehouses"]

#: Other spellings a sheet might use for a column.
_ALIASES = {
    "item description": "description", "name": "description", "item name": "description",
    "item category": "category",
    "valuation": "valuation method",
    "standard cost per unit": "standard cost/unit", "standard cost": "standard cost/unit",
    "cost": "standard cost/unit", "std cost": "standard cost/unit",
    "storage unit": "storage uom", "consumption unit": "consumption uom",
    "bag capacity (kg)": "kg per bag", "bag capacity": "kg per bag", "kg/bag": "kg per bag",
    "hsn": "hsn code", "hsn / sac": "hsn code", "hsn/sac": "hsn code",
    "warehouse": "warehouses", "office": "warehouses", "offices": "warehouses",
}

_REQUIRED = ["description", "category", "valuation method", "standard cost/unit", "usage"]


class ItemImportError(ValueError):
    """A file or a set of rows that cannot be imported; the message says why."""


def _key(text):
    return " ".join(str(text or "").split()).lower()


def _blank(value):
    return value is None or str(value).strip() == ""


def _decimal(value):
    """(Decimal or None, ok): None with ok for a blank, None without ok for junk."""
    if _blank(value):
        return None, True
    try:
        number = Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError):
        return None, False
    return (number, True) if number.is_finite() else (None, False)


def _read(upload):
    name = (getattr(upload, "name", "") or "").lower()
    if name.endswith(".csv"):
        raw = upload.read()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("latin-1")
        table = list(csv.reader(io.StringIO(text)))
    elif name.endswith(".xlsx"):
        import openpyxl
        try:
            workbook = openpyxl.load_workbook(upload, read_only=True, data_only=True)
        except Exception:
            raise ItemImportError("The file could not be read as an Excel workbook.")
        sheet = workbook["Items"] if "Items" in workbook.sheetnames else workbook.active
        table = [list(row) for row in sheet.iter_rows(values_only=True)]
    else:
        raise ItemImportError("Upload an Excel (.xlsx) or CSV (.csv) file.")
    if not table:
        raise ItemImportError("The file is empty.")
    headers = [_ALIASES.get(_key(cell), _key(cell)) for cell in table[0]]
    return headers, table[1:]


def _masters():
    uoms = {}
    for uom in UnitOfMeasurement.objects.all():
        uoms[_key(uom.name)] = uom
        if uom.symbol:
            uoms.setdefault(_key(uom.symbol), uom)
    valuation = {}
    for value, label in Item.VALUATION_METHODS:
        valuation[_key(value)] = value
        valuation[_key(label)] = value
    valuation.setdefault("weighted avg.", "Weighted Average")
    valuation.setdefault("weighted avg", "Weighted Average")
    return {
        "categories": {_key(c.name): c for c in ItemCategory.objects.all()},
        "uoms": uoms,
        "warehouses": {_key(w.name): w for w in Warehouse.objects.all()},
        "valuation": valuation,
        "usage": {"produced": "Produced", "production": "Produced",
                  "sales": "Sales", "sale": "Sales"},
        "existing": {(_key(i.description), i.category_id): i.item_code
                     for i in Item.objects.only("description", "category_id", "item_code")},
    }


def parse_item_upload(upload):
    """Every row of an item sheet, matched to the masters. Saves nothing.

    Each row is "create" or "error" with its problems listed; blank rows are
    skipped. An item whose description and category already exist is an error,
    so a sheet uploaded twice does not make everything twice."""
    headers, body = _read(upload)
    missing = [name for name in _REQUIRED if name not in headers]
    if missing:
        raise ItemImportError("The file is missing these columns: %s." % ", ".join(
            next(c for c in COLUMNS if _key(c) == name) for name in missing))
    index = {name: position for position, name in enumerate(headers)}

    def cell(row, name):
        position = index.get(name)
        return row[position] if position is not None and position < len(row) else None

    masters = _masters()
    rows = []
    seen = set()
    counts = {"create": 0, "error": 0}
    for row_num, raw in enumerate(body, start=2):
        if all(_blank(value) for value in raw):
            continue
        problems = []
        description = " ".join(str(cell(raw, "description") or "").split())
        category_text = str(cell(raw, "category") or "").strip()
        category = masters["categories"].get(_key(category_text))
        if not description:
            problems.append("Description is missing")
        if not category_text:
            problems.append("Category is missing")
        elif category is None:
            problems.append("No category named %s" % category_text)

        valuation = masters["valuation"].get(_key(cell(raw, "valuation method")))
        if valuation is None:
            problems.append("Valuation method not recognised (Weighted Average, Standard Costing, FIFO, LIFO or FEFO)")

        cost, cost_ok = _decimal(cell(raw, "standard cost/unit"))
        if not cost_ok or cost is None:
            problems.append("Standard cost is missing or not a number")
        elif cost < 0:
            problems.append("Standard cost cannot be negative")

        usage = masters["usage"].get(_key(cell(raw, "usage")))
        if usage is None:
            problems.append("Usage must be Produced or Sales")

        units = {}
        for field in ("storage uom", "consumption uom"):
            text = str(cell(raw, field) or "").strip()
            unit = masters["uoms"].get(_key(text)) if text else None
            if text and unit is None:
                problems.append("No unit named %s" % text)
            units[field] = unit

        kg_per_bag, kg_ok = _decimal(cell(raw, "kg per bag"))
        if not kg_ok or (kg_per_bag is not None and kg_per_bag <= 0):
            problems.append("Kg per bag must be a number above zero")

        warehouse_text = str(cell(raw, "warehouses") or "").strip()
        warehouses = []
        if _key(warehouse_text) == "all":
            warehouses = list(masters["warehouses"].values())
        elif warehouse_text:
            for name in warehouse_text.split(","):
                if not name.strip():
                    continue
                warehouse = masters["warehouses"].get(_key(name))
                if warehouse is None:
                    problems.append("No warehouse named %s" % name.strip())
                else:
                    warehouses.append(warehouse)

        if description and category is not None:
            key = (_key(description), category.id)
            if key in seen:
                problems.append("The same item appears earlier in the file")
            elif key in masters["existing"]:
                problems.append("Already exists as %s" % masters["existing"][key])
            seen.add(key)

        action = "error" if problems else "create"
        counts[action] += 1
        rows.append({
            "row": row_num,
            "description": description,
            "category": category.id if category else None,
            "category_name": category.name if category else category_text,
            "valuation_method": valuation,
            "standard_cost_per_unit": str(cost) if cost is not None and cost_ok else None,
            "usage": usage,
            "storage_uom": units["storage uom"].id if units["storage uom"] else None,
            "storage_uom_name": units["storage uom"].name if units["storage uom"] else "",
            "consumption_uom": units["consumption uom"].id if units["consumption uom"] else None,
            "consumption_uom_name": units["consumption uom"].name if units["consumption uom"] else "",
            "kg_per_bag": str(kg_per_bag) if kg_per_bag is not None and kg_ok else None,
            "hsn_code": str(cell(raw, "hsn code") or "").strip()[:100],
            "warehouses": [w.id for w in warehouses],
            "warehouse_names": [w.name for w in warehouses],
            "action": action,
            "message": "; ".join(problems),
        })
    if not rows:
        raise ItemImportError("No item rows were found in the file.")
    if len(rows) > MAX_ROWS:
        raise ItemImportError("Import at most %d items at a time." % MAX_ROWS)
    return {"rows": rows, "counts": counts}


def import_items(rows):
    """Create previewed items, every one or none.

    Each row is checked again rather than trusted, since the masters or the
    item list may have changed between the preview and the save."""
    if not isinstance(rows, list) or not rows:
        raise ItemImportError("There are no items to import.")
    if len(rows) > MAX_ROWS:
        raise ItemImportError("Import at most %d items at a time." % MAX_ROWS)

    categories = set(ItemCategory.objects.values_list("id", flat=True))
    uoms = set(UnitOfMeasurement.objects.values_list("id", flat=True))
    warehouses = set(Warehouse.objects.values_list("id", flat=True))
    existing = {(_key(d), c) for d, c in Item.objects.values_list("description", "category_id")}
    valuations = {value for value, _label in Item.VALUATION_METHODS}

    cleaned = []
    seen = set()
    for number, row in enumerate(rows, start=1):
        row = row if isinstance(row, dict) else {}
        description = " ".join(str(row.get("description") or "").split())
        category = row.get("category")
        where = "Row %d" % number
        if not description:
            raise ItemImportError("%s: description is missing." % where)
        if category not in categories:
            raise ItemImportError("%s: the category no longer exists. Preview again." % where)
        if row.get("valuation_method") not in valuations:
            raise ItemImportError("%s: valuation method not recognised." % where)
        if row.get("usage") not in ("Produced", "Sales"):
            raise ItemImportError("%s: usage must be Produced or Sales." % where)
        cost, cost_ok = _decimal(row.get("standard_cost_per_unit"))
        if not cost_ok or cost is None or cost < 0:
            raise ItemImportError("%s: standard cost is missing or not a number." % where)
        kg_per_bag, kg_ok = _decimal(row.get("kg_per_bag"))
        if not kg_ok or (kg_per_bag is not None and kg_per_bag <= 0):
            raise ItemImportError("%s: kg per bag must be a number above zero." % where)
        for field in ("storage_uom", "consumption_uom"):
            if row.get(field) is not None and row.get(field) not in uoms:
                raise ItemImportError("%s: a unit no longer exists. Preview again." % where)
        chosen = [w for w in (row.get("warehouses") or [])]
        if any(w not in warehouses for w in chosen):
            raise ItemImportError("%s: a warehouse no longer exists. Preview again." % where)
        key = (_key(description), category)
        if key in seen or key in existing:
            raise ItemImportError("%s: %s already exists in that category." % (where, description))
        seen.add(key)
        cleaned.append({
            "description": description, "category_id": category,
            "valuation_method": row["valuation_method"], "standard_cost_per_unit": cost,
            "usage": row["usage"], "storage_uom_id": row.get("storage_uom"),
            "consumption_uom_id": row.get("consumption_uom"), "kg_per_bag": kg_per_bag,
            "hsn_code": (str(row.get("hsn_code") or "").strip()[:100]) or None,
            "warehouses": chosen,
        })

    created = []
    with transaction.atomic():
        for fields in cleaned:
            chosen = fields.pop("warehouses")
            item = Item.objects.create(**fields)
            if chosen:
                item.warehouse.set(chosen)
            created.append(item)
    return created


def item_template_workbook():
    """An .xlsx with the item columns to fill in, and a second sheet listing
    the categories, units, warehouses and choices the columns accept."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Items"
    sheet.append(COLUMNS)
    for position, cell in enumerate(sheet[1]):
        cell.font = Font(bold=True)
        if position < len(_REQUIRED):
            cell.fill = PatternFill("solid", fgColor="FFF4CC")
    for column, width in zip("ABCDEFGHIJ", (34, 18, 18, 18, 12, 14, 16, 12, 14, 28)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"

    lists = workbook.create_sheet("Lists")
    lists.append(["Categories", "Units (name or symbol)", "Warehouses", "Valuation Methods", "Usage"])
    for cell in lists[1]:
        cell.font = Font(bold=True)
    columns = [
        list(ItemCategory.objects.order_by("name").values_list("name", flat=True)),
        [u.name + (" (%s)" % u.symbol if u.symbol else "") for u in UnitOfMeasurement.objects.order_by("name")],
        ["All"] + list(Warehouse.objects.order_by("name").values_list("name", flat=True)),
        [value for value, _label in Item.VALUATION_METHODS],
        ["Produced", "Sales"],
    ]
    for position in range(max(len(c) for c in columns)):
        lists.append([c[position] if position < len(c) else None for c in columns])
    for column in "ABCDE":
        lists.column_dimensions[column].width = 24

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
