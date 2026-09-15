"""Item Price List — the current-price view, price history, bulk revisions and
uploads.

The list is a dated history: an item's price on a day is its most recent entry
effective on or before that day, which is what transfers are valued at
(inventory.services.pricing.item_issue_price). Everything here reads the
history the same way, so the price this page calls Active is the price a
transfer made today would take.
"""
import csv
import io
from collections import defaultdict
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation

from django.db import transaction
from django.db.models import Max, Q
from django.utils import timezone

from inventory.models import Item, ItemPriceList, ItemPriceListAudit

TWO_PLACES = Decimal("0.01")
MAX_ROWS = 5000

STATUS_ACTIVE = "active"
STATUS_UPCOMING = "upcoming"
STATUS_NOT_PRICED = "not_priced"
STATUS_SUPERSEDED = "superseded"
STATUS_LABELS = {
    STATUS_ACTIVE: "Active",
    STATUS_UPCOMING: "Upcoming",
    STATUS_NOT_PRICED: "Not Priced",
    STATUS_SUPERSEDED: "Superseded",
}


class PriceRowError(ValueError):
    """A revision or upload that cannot be used as asked; the message says why."""


# --- small readers -----------------------------------------------------------

def _money(value):
    return Decimal(value).quantize(TWO_PLACES, rounding=ROUND_HALF_UP)


def _text(value):
    return None if value is None else str(_money(value))


def _pct(new, old):
    if new is None or not old:
        return None
    change = (Decimal(new) - Decimal(old)) / Decimal(old) * 100
    return str(change.quantize(TWO_PLACES, rounding=ROUND_HALF_UP))


def _iso(value):
    return value.isoformat() if value else None


def _unit(item):
    uom = item.storage_uom or item.consumption_uom
    if not uom:
        return ""
    return uom.symbol or uom.name


def parse_date(value):
    """A date from a form, JSON or a spreadsheet cell: a date, an ISO string,
    or DD-MM-YYYY / DD.MM.YYYY / DD/MM/YYYY."""
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def parse_price(value):
    """A price rounded to paise, None when blank. Raises PriceRowError when
    the value is not a number."""
    if value is None or str(value).strip() == "":
        return None
    try:
        price = Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError):
        raise PriceRowError("not a number")
    if not price.is_finite():
        raise PriceRowError("not a number")
    return _money(price)


def _clean_ids(values):
    ids = []
    for value in values or []:
        text = str(value).strip()
        if text.isdigit() and int(text) not in ids:
            ids.append(int(text))
    return ids


def _entries_by_item(item_ids):
    """Each item's price entries, newest effective date first."""
    grouped = defaultdict(list)
    entries = (ItemPriceList.objects.filter(item_id__in=item_ids)
               .order_by("-effective_date", "-id"))
    for entry in entries:
        grouped[entry.item_id].append(entry)
    return grouped


def _in_force(entries, on_date):
    return next((e for e in entries if e.effective_date <= on_date), None)


def _item_info(item):
    return {
        "item": item.id,
        "item_code": item.item_code,
        "item_name": item.description,
        "category": item.category.name if item.category_id else "",
        "category_id": item.category_id,
        "item_type": item.type,
        "unit": _unit(item),
    }


# --- the list ----------------------------------------------------------------

def price_overview(today=None, category=None, item_type=None, status=None, search=None):
    """Every item once, with the price in force today, the one before it and
    any price already set for a later date.

    Active: priced today. Upcoming: only priced from a later date. Not Priced:
    no price at all, so a transfer of it would be refused."""
    today = today or timezone.localdate()
    items = (Item.objects.select_related("category", "storage_uom", "consumption_uom")
             .order_by("item_code"))
    if category and str(category).isdigit():
        items = items.filter(category_id=int(category))
    if item_type:
        items = items.filter(type=item_type)
    if search and search.strip():
        term = search.strip()
        items = items.filter(Q(item_code__icontains=term) | Q(description__icontains=term))
    items = list(items)
    grouped = _entries_by_item([item.id for item in items])

    rows = []
    for item in items:
        entries = grouped.get(item.id, [])
        past = [e for e in entries if e.effective_date <= today]
        future = [e for e in entries if e.effective_date > today]
        current = past[0] if past else None
        previous = past[1] if len(past) > 1 else None
        upcoming = future[-1] if future else None

        if current:
            state, shown = STATUS_ACTIVE, current
        elif upcoming:
            state, shown = STATUS_UPCOMING, upcoming
        else:
            state, shown = STATUS_NOT_PRICED, None

        row = _item_info(item)
        row.update({
            "entry": shown.id if shown else None,
            "price": _text(shown.price) if shown else None,
            "effective_date": _iso(shown.effective_date) if shown else None,
            "previous_price": _text(previous.price) if current and previous else None,
            "previous_date": _iso(previous.effective_date) if current and previous else None,
            "change_pct": _pct(current.price, previous.price) if current and previous else None,
            "next_price": _text(upcoming.price) if current and upcoming else None,
            "next_date": _iso(upcoming.effective_date) if current and upcoming else None,
            "status": state,
            "status_label": STATUS_LABELS[state],
            "entries": len(entries),
        })
        rows.append(row)

    if status and status != "all":
        rows = [row for row in rows if row["status"] == status]
    return rows


def last_price_change():
    """When a price was last created, edited or deleted, in local time."""
    stamp = ItemPriceListAudit.objects.aggregate(last=Max("created_at"))["last"]
    return timezone.localtime(stamp) if stamp else None


def audit_rows(item_id=None, limit=500):
    logs = ItemPriceListAudit.objects.select_related("user").order_by("-created_at", "-id")
    if item_id and str(item_id).isdigit():
        logs = logs.filter(item_ref=int(item_id))
    rows = []
    for log in logs[:limit]:
        by = log.user_label
        if not by and log.user_id:
            by = log.user.get_full_name() or log.user.get_username()
        rows.append({
            "id": log.id,
            "when": timezone.localtime(log.created_at).strftime("%d %b %Y %I:%M %p"),
            "item": log.item_label,
            "item_ref": log.item_ref,
            "action": log.action,
            "action_label": log.get_action_display(),
            "old_price": _text(log.old_price),
            "new_price": _text(log.new_price),
            "old_date": _iso(log.old_effective_date),
            "new_date": _iso(log.new_effective_date),
            "by": by or "System",
            "source": log.source,
            "note": log.note,
        })
    return rows


def item_price_history(item, today=None):
    """All of one item's prices, newest first, each marked Upcoming, Active or
    Superseded, with its change against the price before it."""
    today = today or timezone.localdate()
    entries = list(item.price_list_entries.order_by("-effective_date", "-id"))
    current = _in_force(entries, today)
    rows = []
    for index, entry in enumerate(entries):
        older = entries[index + 1] if index + 1 < len(entries) else None
        if entry.effective_date > today:
            state = STATUS_UPCOMING
        elif current and entry.id == current.id:
            state = STATUS_ACTIVE
        else:
            state = STATUS_SUPERSEDED
        rows.append({
            "id": entry.id,
            "price": _text(entry.price),
            "effective_date": _iso(entry.effective_date),
            "change_pct": _pct(entry.price, older.price) if older else None,
            "status": state,
            "status_label": STATUS_LABELS[state],
        })
    return {
        "item": _item_info(item),
        "today": today.isoformat(),
        "entries": rows,
        "audit": audit_rows(item_id=item.id, limit=200),
    }


# --- bulk revision -----------------------------------------------------------

def revise_preview(item_ids, mode, value, effective_date):
    """New prices for a revision by a percentage or a fixed amount, worked
    from each item's price in force on the new effective date. Saves nothing.

    Each row is "create" (a new dated price), "replace" (the item already has
    a price on that date) or "skip" with the reason."""
    if mode not in ("percent", "amount"):
        raise PriceRowError("Choose whether to revise by a percentage or a fixed amount.")
    try:
        value = Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError):
        raise PriceRowError("Enter the revision as a number.")
    if not value.is_finite() or value == 0:
        raise PriceRowError("Enter a revision other than zero.")
    if mode == "percent" and value <= -100:
        raise PriceRowError("A decrease must be less than 100%.")
    on_date = parse_date(effective_date)
    if not on_date:
        raise PriceRowError("Choose the effective date for the new prices.")
    ids = _clean_ids(item_ids)
    if not ids:
        raise PriceRowError("Select at least one item to revise.")
    if len(ids) > MAX_ROWS:
        raise PriceRowError("Revise at most %d items at a time." % MAX_ROWS)

    items = list(Item.objects.filter(id__in=ids).select_related(
        "category", "storage_uom", "consumption_uom").order_by("item_code"))
    grouped = _entries_by_item(ids)
    rows = []
    counts = {"create": 0, "replace": 0, "skip": 0}
    for item in items:
        base = _in_force(grouped.get(item.id, []), on_date)
        row = _item_info(item)
        row.update({
            "current_price": _text(base.price) if base else None,
            "current_date": _iso(base.effective_date) if base else None,
            "new_price": None,
            "change_pct": None,
            "effective_date": on_date.isoformat(),
            "action": "skip",
            "message": "",
        })
        if base is None:
            row["message"] = "No price in force on %s to revise" % on_date.strftime("%d.%m.%Y")
        else:
            if mode == "percent":
                new = _money(base.price * (1 + value / 100))
            else:
                new = _money(base.price + value)
            if new <= 0:
                row["message"] = "Revised price would be zero or less"
            else:
                row["new_price"] = str(new)
                row["change_pct"] = _pct(new, base.price)
                if base.effective_date == on_date:
                    row["action"] = "replace"
                    row["message"] = "Replaces the price already set for this date"
                else:
                    row["action"] = "create"
        counts[row["action"]] += 1
        rows.append(row)
    return {"rows": rows, "counts": counts}


# --- upload ------------------------------------------------------------------

_CODE_HEADERS = ("item code", "item_code", "code")
_PRICE_HEADERS = ("new price", "new_price", "price", "rate")
_DATE_HEADERS = ("effective date", "effective_date", "date")


def _read_upload(upload):
    name = (getattr(upload, "name", "") or "").lower()
    if name.endswith(".csv"):
        raw = upload.read()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("latin-1")
        table = [row for row in csv.reader(io.StringIO(text))]
    elif name.endswith(".xlsx"):
        import openpyxl
        try:
            workbook = openpyxl.load_workbook(upload, read_only=True, data_only=True)
        except Exception:
            raise PriceRowError("The file could not be read as an Excel workbook.")
        table = [list(row) for row in workbook.active.iter_rows(values_only=True)]
    else:
        raise PriceRowError("Upload an Excel (.xlsx) or CSV (.csv) file.")
    if not table:
        raise PriceRowError("The file is empty.")
    headers = [str(cell or "").strip().lower() for cell in table[0]]
    return headers, table[1:]


def _column(headers, names):
    for name in names:
        if name in headers:
            return headers.index(name)
    return None


def _cell(row, col):
    if col is None or col >= len(row):
        return None
    return row[col]


def _blank(value):
    return value is None or str(value).strip() == ""


def parse_price_upload(upload, default_date=None):
    """Rows of an uploaded price file matched to items by Item Code. Saves
    nothing.

    Rows with a blank New Price are left out: the template lists every item,
    and only the ones filled in are being changed. A row without its own
    Effective Date takes default_date."""
    headers, body = _read_upload(upload)
    code_col = _column(headers, _CODE_HEADERS)
    price_col = _column(headers, _PRICE_HEADERS)
    date_col = _column(headers, _DATE_HEADERS)
    if code_col is None or price_col is None:
        raise PriceRowError("The file needs an Item Code column and a New Price (or Price) column.")
    fallback = parse_date(default_date)

    pending = []
    for row_num, row in enumerate(body, start=2):
        if _blank(_cell(row, price_col)):
            continue
        pending.append((row_num, str(_cell(row, code_col) or "").strip(),
                        _cell(row, price_col), _cell(row, date_col)))
    if not pending:
        raise PriceRowError("No rows with a new price were found in the file.")
    if len(pending) > MAX_ROWS:
        raise PriceRowError("Upload at most %d prices at a time." % MAX_ROWS)

    items = {item.item_code.upper(): item for item in Item.objects.all()}
    grouped = _entries_by_item([item.id for item in items.values()])
    seen = set()
    rows = []
    counts = {"create": 0, "replace": 0, "unchanged": 0, "error": 0}
    for row_num, code, raw_price, raw_date in pending:
        item = items.get(code.upper()) if code else None
        row = {
            "row": row_num, "item_code": code,
            "item": item.id if item else None,
            "item_name": item.description if item else "",
            "current_price": None, "new_price": None, "change_pct": None,
            "effective_date": None, "action": "error", "message": "",
        }
        try:
            price = parse_price(raw_price)
            price_ok = True
        except PriceRowError:
            price, price_ok = None, False
        on_date = fallback if _blank(raw_date) else parse_date(raw_date)

        if not code:
            row["message"] = "Item code is missing"
        elif item is None:
            row["message"] = "No item with code %s" % code
        elif not price_ok:
            row["message"] = "Price is not a number"
        elif price <= 0:
            row["message"] = "Price must be more than zero"
        elif not _blank(raw_date) and on_date is None:
            row["message"] = "Effective date not understood (use YYYY-MM-DD or DD-MM-YYYY)"
        elif on_date is None:
            row["message"] = "No effective date: fill the column or choose a date"
        elif (item.id, on_date) in seen:
            row["message"] = "Same item and date appear earlier in the file"
        else:
            seen.add((item.id, on_date))
            base = _in_force(grouped.get(item.id, []), on_date)
            row.update({
                "current_price": _text(base.price) if base else None,
                "new_price": str(price),
                "change_pct": _pct(price, base.price) if base else None,
                "effective_date": on_date.isoformat(),
            })
            if base and base.effective_date == on_date:
                if base.price == price:
                    row["action"] = "unchanged"
                    row["message"] = "Same as the price already set for this date"
                else:
                    row["action"] = "replace"
                    row["message"] = "Replaces the price already set for this date"
            else:
                row["action"] = "create"
        counts[row["action"]] += 1
        rows.append(row)
    return {"rows": rows, "counts": counts}


# --- saving ------------------------------------------------------------------

def apply_price_rows(rows, *, source, note="", can_replace=True):
    """Save dated prices, all of them or none.

    rows: [{"item": id, "price": "45.15", "effective_date": "YYYY-MM-DD"}].
    A new date adds an entry; a date the item is already priced on replaces
    that price, which needs can_replace (Edit rights)."""
    from inventory.price_audit import price_change

    if not isinstance(rows, list) or not rows:
        raise PriceRowError("There are no prices to save.")
    if len(rows) > MAX_ROWS:
        raise PriceRowError("Save at most %d prices at a time." % MAX_ROWS)

    cleaned = []
    seen = set()
    for index, row in enumerate(rows, start=1):
        row = row if isinstance(row, dict) else {}
        item_id = str(row.get("item") or "").strip()
        if not item_id.isdigit():
            raise PriceRowError("Row %d: item is missing." % index)
        try:
            price = parse_price(row.get("price"))
        except PriceRowError:
            raise PriceRowError("Row %d: price is not a number." % index)
        if price is None or price <= 0:
            raise PriceRowError("Row %d: price must be more than zero." % index)
        on_date = parse_date(row.get("effective_date"))
        if on_date is None:
            raise PriceRowError("Row %d: effective date is missing." % index)
        key = (int(item_id), on_date)
        if key in seen:
            raise PriceRowError("Row %d: the same item and date appear twice." % index)
        seen.add(key)
        cleaned.append((int(item_id), on_date, price))

    item_ids = {item_id for item_id, _, _ in cleaned}
    found = set(Item.objects.filter(id__in=item_ids).values_list("id", flat=True))
    if item_ids - found:
        raise PriceRowError("One or more items no longer exist. Preview again.")

    existing = {
        (entry.item_id, entry.effective_date): entry
        for entry in ItemPriceList.objects.filter(
            item_id__in=item_ids, effective_date__in={d for _, d, _ in cleaned})
    }
    replacing = [(i, d, p) for i, d, p in cleaned
                 if (i, d) in existing and existing[(i, d)].price != p]
    if replacing and not can_replace:
        raise PriceRowError(
            "%d item%s already priced on that date. Changing an existing price "
            "needs Edit rights on the Item Price List."
            % (len(replacing), " is" if len(replacing) == 1 else "s are"))

    created = updated = unchanged = 0
    with transaction.atomic(), price_change(source, note):
        for item_id, on_date, price in cleaned:
            entry = existing.get((item_id, on_date))
            if entry is None:
                ItemPriceList.objects.create(item_id=item_id, price=price, effective_date=on_date)
                created += 1
            elif entry.price != price:
                entry.price = price
                entry.save(update_fields=["price"])
                updated += 1
            else:
                unchanged += 1
    return {"created": created, "updated": updated, "unchanged": unchanged}


def price_template_workbook(today=None):
    """An .xlsx of every item with its current price and an empty New Price
    column, to fill in and upload."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Prices"
    headers = ["Item Code", "Item Name", "Category", "Unit", "Current Price",
               "Current Since", "New Price", "Effective Date"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for col in ("G", "H"):
        sheet["%s1" % col].fill = PatternFill("solid", fgColor="FFF4CC")

    for row in price_overview(today=today):
        active = row["status"] == STATUS_ACTIVE
        sheet.append([
            row["item_code"], row["item_name"], row["category"], row["unit"],
            float(row["price"]) if active else None,
            row["effective_date"] if active else None,
            None, None,
        ])
    for col, width in zip("ABCDEFGH", (14, 36, 18, 8, 14, 14, 12, 15)):
        sheet.column_dimensions[col].width = width
    sheet.freeze_panes = "A2"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
