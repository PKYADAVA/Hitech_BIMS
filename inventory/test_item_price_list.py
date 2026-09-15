"""Item Price List: the current-price view, per-item history, bulk revision,
upload, and the change log kept of every price created, edited or deleted.

The list is a dated history, and a transfer takes the price in force on its
date. So "Active" here has to mean exactly that price, a revision has to work
from it, and a change to it has to leave a record of who made it.
"""
import io
import json
from datetime import timedelta
from decimal import Decimal
from types import SimpleNamespace

import openpyxl
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from alerts.context import clear_current_request, set_current_request
from inventory.models import (Item, ItemCategory, ItemPriceList,
                              ItemPriceListAudit, UnitOfMeasurement)
from inventory.price_audit import SOURCE_API, SOURCE_BULK, SOURCE_PAGE
from inventory.services.price_list import (PriceRowError, apply_price_rows,
                                           item_price_history,
                                           parse_price_upload, price_overview,
                                           revise_preview)


class PriceListBase(TestCase):

    def setUp(self):
        self.today = timezone.localdate()
        self.feed = ItemCategory.objects.create(name="Feed")
        self.medicine = ItemCategory.objects.create(name="Medicine")
        kg = UnitOfMeasurement.objects.create(name="Kilogram", symbol="Kg")
        common = dict(valuation_method="Weighted Average", usage="Produced",
                      source="Purchased", item_account="Expense",
                      standard_cost_per_unit=0)
        self.starter = Item.objects.create(
            description="Broiler Starter Feed", category=self.feed,
            type="Raw Material", storage_uom=kg, **common)
        self.vaccine = Item.objects.create(
            description="Gumboro Vaccine", category=self.medicine,
            type="Finished Goods", **common)
        self.tonic = Item.objects.create(
            description="Liver Tonic", category=self.medicine,
            type="Finished Goods", **common)
        self.user = get_user_model().objects.create_superuser(
            "pricer", "p@x.com", "Str0ngPass!")
        self.client.force_login(self.user)

    def day(self, n):
        return self.today + timedelta(days=n)

    def price(self, item, amount, days=0):
        return ItemPriceList.objects.create(
            item=item, price=Decimal(str(amount)), effective_date=self.day(days))

    def row_for(self, rows, item):
        return next(r for r in rows if r["item"] == item.id)


class OverviewTests(PriceListBase):

    def test_the_price_in_force_today_is_active_beside_the_one_before_it(self):
        self.price(self.starter, 40, -30)
        self.price(self.starter, 42, -5)
        row = self.row_for(price_overview(today=self.today), self.starter)
        self.assertEqual(row["status"], "active")
        self.assertEqual(row["price"], "42.00")
        self.assertEqual(row["effective_date"], self.day(-5).isoformat())
        self.assertEqual(row["previous_price"], "40.00")
        self.assertEqual(row["change_pct"], "5.00")
        self.assertEqual(row["unit"], "Kg")

    def test_a_price_set_for_a_later_date_is_shown_as_coming_not_as_current(self):
        """A transfer today still takes 42, so 42 is the price shown."""
        self.price(self.starter, 42, -5)
        self.price(self.starter, 45, 10)
        row = self.row_for(price_overview(today=self.today), self.starter)
        self.assertEqual(row["price"], "42.00")
        self.assertEqual(row["next_price"], "45.00")
        self.assertEqual(row["next_date"], self.day(10).isoformat())

    def test_an_item_priced_only_from_a_later_date_is_upcoming(self):
        self.price(self.vaccine, 5, 3)
        row = self.row_for(price_overview(today=self.today), self.vaccine)
        self.assertEqual(row["status"], "upcoming")
        self.assertEqual(row["price"], "5.00")
        self.assertIsNone(row["previous_price"])

    def test_an_item_never_priced_is_listed_as_not_priced(self):
        """Worth seeing: a transfer of it would be refused."""
        row = self.row_for(price_overview(today=self.today), self.tonic)
        self.assertEqual(row["status"], "not_priced")
        self.assertEqual(row["status_label"], "Not Priced")
        self.assertIsNone(row["price"])
        self.assertIsNone(row["entry"])

    def test_the_filters(self):
        self.price(self.starter, 42, -1)
        self.price(self.vaccine, 5, 2)

        def ids(**filters):
            return {r["item"] for r in price_overview(today=self.today, **filters)}

        self.assertEqual(ids(category=self.medicine.id), {self.vaccine.id, self.tonic.id})
        self.assertEqual(ids(status="upcoming"), {self.vaccine.id})
        self.assertEqual(ids(status="not_priced"), {self.tonic.id})
        self.assertEqual(ids(search="gumboro"), {self.vaccine.id})
        self.assertEqual(ids(search=self.starter.item_code.lower()), {self.starter.id})

    def test_the_page_data_counts_every_status_before_the_status_filter(self):
        self.price(self.starter, 42, -1)
        self.price(self.vaccine, 5, 2)
        res = self.client.get(reverse("item_price_list_overview_data"),
                              {"status": "active"}).json()
        self.assertEqual([r["item"] for r in res["rows"]], [self.starter.id])
        self.assertEqual(res["counts"], {"active": 1, "upcoming": 1, "not_priced": 1,
                                         "inactive": 0, "all": 3})
        self.assertTrue(res["last_updated"])

    def test_a_new_item_shows_up_at_once_as_not_priced(self):
        """Nothing to do on the price list: an item created on the Items page
        is listed straight away, waiting for its price."""
        response = self.client.post(reverse("item_create"), json.dumps({
            "description": "Grower Feed", "category": self.feed.id,
            "valuation_method": "Weighted Average", "standard_cost_per_unit": "0",
            "usage": "Produced",
        }), content_type="application/json")
        self.assertEqual(response.status_code, 201, response.content)
        grower = Item.objects.get(description="Grower Feed")

        rows = self.client.get(reverse("item_price_list_overview_data")).json()["rows"]
        row = self.row_for(rows, grower)
        self.assertEqual((row["status"], row["price"], row["is_active"]), ("not_priced", None, True))
        page = self.client.get(reverse("item_price_list"))
        self.assertIn(grower.id, [i["id"] for i in page.context["items_json"]])

    def test_the_page_renders(self):
        response = self.client.get(reverse("item_price_list"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'id="ipl-table"')
        self.assertContains(response, 'id="ipl-items"')


class HistoryTests(PriceListBase):

    def test_each_price_is_marked_superseded_active_or_upcoming(self):
        old = self.price(self.starter, 40, -30)
        now = self.price(self.starter, 42, -5)
        later = self.price(self.starter, 45, 10)
        res = item_price_history(self.starter, today=self.today)
        self.assertEqual([(e["id"], e["status"]) for e in res["entries"]],
                         [(later.id, "upcoming"), (now.id, "active"), (old.id, "superseded")])
        self.assertEqual(res["entries"][1]["change_pct"], "5.00")
        self.assertIsNone(res["entries"][2]["change_pct"])

    def test_the_history_endpoint_carries_the_prices_and_the_log(self):
        self.price(self.starter, 42, -5)
        res = self.client.get(reverse("item_price_list_history_data",
                                      args=[self.starter.id])).json()
        self.assertEqual(res["item"]["item_code"], self.starter.item_code)
        self.assertEqual(len(res["entries"]), 1)
        self.assertEqual([a["action"] for a in res["audit"]], ["create"])


class ChangeLogTests(PriceListBase):

    def test_adding_a_price_on_the_page_records_who_and_where(self):
        self.client.post(reverse("item_price_list_create"), json.dumps({
            "effective_date": self.today.isoformat(),
            "rows": [{"item": self.starter.id, "price": "42.50"}],
        }), content_type="application/json")
        log = ItemPriceListAudit.objects.get()
        self.assertEqual(log.action, "create")
        self.assertIsNone(log.old_price)
        self.assertEqual(log.new_price, Decimal("42.50"))
        self.assertEqual(log.new_effective_date, self.today)
        self.assertEqual(log.user, self.user)
        self.assertEqual(log.source, SOURCE_PAGE)
        self.assertIn("Broiler Starter Feed", log.item_label)

    def test_editing_records_the_price_and_date_before_and_after(self):
        entry = self.price(self.starter, 42, -5)
        ItemPriceListAudit.objects.all().delete()
        self.client.put(reverse("item_price_list_update", args=[entry.id]), json.dumps({
            "price": "44", "effective_date": self.day(-2).isoformat(),
        }), content_type="application/json")
        log = ItemPriceListAudit.objects.get()
        self.assertEqual(log.action, "update")
        self.assertEqual((log.old_price, log.new_price), (Decimal("42.00"), Decimal("44.00")))
        self.assertEqual((log.old_effective_date, log.new_effective_date),
                         (self.day(-5), self.day(-2)))
        self.assertEqual(log.user, self.user)

    def test_a_save_that_changes_nothing_is_not_recorded(self):
        entry = self.price(self.starter, 42, -5)
        ItemPriceListAudit.objects.all().delete()
        entry.save()
        self.assertFalse(ItemPriceListAudit.objects.exists())

    def test_deleting_keeps_a_record_of_what_was_deleted(self):
        entry = self.price(self.starter, 42, -5)
        ItemPriceListAudit.objects.all().delete()
        self.client.delete(reverse("item_price_list_delete", args=[entry.id]))
        log = ItemPriceListAudit.objects.get()
        self.assertEqual(log.action, "delete")
        self.assertEqual(log.old_price, Decimal("42.00"))
        self.assertIsNone(log.new_price)
        self.assertEqual(log.price_entry_ref, entry.id)

    def test_deleting_the_item_leaves_its_price_log_behind(self):
        self.price(self.tonic, 320, -1)
        self.tonic.delete()
        actions = sorted(ItemPriceListAudit.objects.filter(item_label__contains="Liver Tonic")
                         .values_list("action", flat=True))
        self.assertEqual(actions, ["create", "delete"])

    def test_a_change_through_the_mobile_api_is_named_as_such(self):
        set_current_request(SimpleNamespace(path="/api/v1/inventory/price-list/", user=self.user))
        try:
            self.price(self.vaccine, 5)
        finally:
            clear_current_request()
        log = ItemPriceListAudit.objects.latest("id")
        self.assertEqual(log.source, SOURCE_API)
        self.assertEqual(log.user_id, self.user.id)

    def test_the_change_log_endpoint_filters_by_item(self):
        self.price(self.starter, 42, -1)
        self.price(self.vaccine, 5, -1)
        rows = self.client.get(reverse("item_price_list_audit_data"),
                               {"item": self.starter.id}).json()["rows"]
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["action"], "create")
        self.assertEqual(rows[0]["new_price"], "42.00")


class BulkRevisionTests(PriceListBase):

    def test_a_percentage_rise_is_rounded_to_paise(self):
        self.price(self.starter, 42, -5)
        row = revise_preview([self.starter.id], "percent", "7.5", self.day(1))["rows"][0]
        self.assertEqual((row["current_price"], row["new_price"]), ("42.00", "45.15"))
        self.assertEqual(row["change_pct"], "7.50")
        self.assertEqual(row["action"], "create")

    def test_a_fixed_amount_cut(self):
        self.price(self.vaccine, 5, -1)
        row = revise_preview([self.vaccine.id], "amount", "-0.75", self.day(1))["rows"][0]
        self.assertEqual(row["new_price"], "4.25")

    def test_it_works_from_the_price_in_force_on_the_new_date(self):
        """Not from a later price already set: that one is still to come."""
        self.price(self.starter, 40, -30)
        self.price(self.starter, 50, 20)
        row = revise_preview([self.starter.id], "percent", 10, self.day(1))["rows"][0]
        self.assertEqual((row["current_price"], row["new_price"]), ("40.00", "44.00"))

    def test_an_item_with_no_price_on_that_date_is_skipped(self):
        row = revise_preview([self.tonic.id], "percent", 10, self.day(1))["rows"][0]
        self.assertEqual(row["action"], "skip")
        self.assertIn("No price in force", row["message"])

    def test_a_cut_to_nothing_is_skipped(self):
        self.price(self.vaccine, 5, -1)
        row = revise_preview([self.vaccine.id], "amount", "-5", self.day(1))["rows"][0]
        self.assertEqual(row["action"], "skip")
        self.assertIn("zero or less", row["message"])

    def test_a_date_already_priced_is_marked_as_a_replacement(self):
        self.price(self.starter, 42, 0)
        row = revise_preview([self.starter.id], "percent", 10, self.today)["rows"][0]
        self.assertEqual(row["action"], "replace")

    def test_unusable_requests_are_refused(self):
        ids = [self.starter.id]
        for args in ((ids, "double", 5, self.today), (ids, "percent", 0, self.today),
                     (ids, "percent", -100, self.today), (ids, "percent", 5, ""),
                     ([], "percent", 5, self.today)):
            with self.assertRaises(PriceRowError):
                revise_preview(*args)

    def test_saving_adds_the_prices_and_logs_them_as_a_bulk_revision(self):
        self.price(self.starter, 42, -5)
        self.price(self.vaccine, 5, -5)
        preview = revise_preview([self.starter.id, self.vaccine.id], "percent", 10, self.day(1))
        rows = [{"item": r["item"], "price": r["new_price"], "effective_date": r["effective_date"]}
                for r in preview["rows"]]
        response = self.client.post(reverse("item_price_list_bulk_create"), json.dumps({
            "rows": rows, "source": "bulk", "note": "Q3 feed revision",
        }), content_type="application/json")
        self.assertEqual(response.status_code, 201, response.content)
        self.assertEqual(response.json()["created"], 2)
        self.assertEqual(ItemPriceList.objects.get(item=self.starter, effective_date=self.day(1)).price,
                         Decimal("46.20"))
        logs = ItemPriceListAudit.objects.filter(new_effective_date=self.day(1))
        self.assertEqual(logs.count(), 2)
        self.assertEqual({(l.source, l.note, l.user_id) for l in logs},
                         {(SOURCE_BULK, "Q3 feed revision", self.user.id)})

    def test_saving_is_all_or_nothing(self):
        rows = [{"item": self.starter.id, "price": "46", "effective_date": self.day(1).isoformat()},
                {"item": self.vaccine.id, "price": "-1", "effective_date": self.day(1).isoformat()}]
        with self.assertRaises(PriceRowError):
            apply_price_rows(rows, source=SOURCE_BULK)
        self.assertFalse(ItemPriceList.objects.filter(effective_date=self.day(1)).exists())

    def test_changing_a_price_already_set_needs_edit_rights(self):
        self.price(self.starter, 42, 1)
        rows = [{"item": self.starter.id, "price": "43", "effective_date": self.day(1).isoformat()}]
        with self.assertRaises(PriceRowError) as refused:
            apply_price_rows(rows, source=SOURCE_BULK, can_replace=False)
        self.assertIn("Edit rights", str(refused.exception))
        self.assertEqual(apply_price_rows(rows, source=SOURCE_BULK, can_replace=True)["updated"], 1)

    def test_adding_a_new_date_does_not_need_edit_rights(self):
        self.price(self.starter, 42, -1)
        rows = [{"item": self.starter.id, "price": "43", "effective_date": self.day(1).isoformat()}]
        self.assertEqual(apply_price_rows(rows, source=SOURCE_BULK, can_replace=False)["created"], 1)


class UploadTests(PriceListBase):

    def workbook(self, rows, headers=("Item Code", "New Price", "Effective Date")):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.append(list(headers))
        for row in rows:
            sheet.append(list(row))
        buffer = io.BytesIO()
        book.save(buffer)
        return SimpleUploadedFile("prices.xlsx", buffer.getvalue())

    def test_rows_are_matched_by_item_code(self):
        self.price(self.starter, 42, -5)
        res = parse_price_upload(self.workbook([
            (self.starter.item_code, 44, None),
            (self.vaccine.item_code.lower(), "5.5", "01-01-2030"),
        ]), default_date=self.day(1).isoformat())
        first, second = res["rows"]
        self.assertEqual((first["item"], first["action"]), (self.starter.id, "create"))
        self.assertEqual((first["current_price"], first["new_price"]), ("42.00", "44.00"))
        self.assertEqual(first["effective_date"], self.day(1).isoformat())
        self.assertEqual((second["item"], second["effective_date"]), (self.vaccine.id, "2030-01-01"))

    def test_a_csv_is_read_the_same_way(self):
        text = "Item Code,Price\n%s,43\n" % self.starter.item_code
        res = parse_price_upload(SimpleUploadedFile("prices.csv", text.encode()),
                                 default_date=self.day(1).isoformat())
        self.assertEqual(res["rows"][0]["action"], "create")

    def test_problems_are_reported_on_their_own_rows(self):
        code, other = self.starter.item_code, self.vaccine.item_code
        res = parse_price_upload(self.workbook([
            ("ITM-9999", 10, None),
            (code, "abc", None),
            (code, 0, None),
            (other, 5, "31-31-2026"),
            (other, 5, None),
            (other, 6, None),
        ]), default_date=self.day(1).isoformat())
        self.assertEqual([r["action"] for r in res["rows"]],
                         ["error", "error", "error", "error", "create", "error"])
        messages = [r["message"] for r in res["rows"]]
        self.assertIn("No item with code ITM-9999", messages[0])
        self.assertIn("not a number", messages[1])
        self.assertIn("more than zero", messages[2])
        self.assertIn("not understood", messages[3])
        self.assertIn("earlier in the file", messages[5])
        self.assertEqual(res["counts"]["error"], 5)

    def test_rows_left_blank_in_the_template_are_left_alone(self):
        res = parse_price_upload(self.workbook([
            (self.starter.item_code, None, None),
            (self.vaccine.item_code, 5, None),
        ]), default_date=self.day(1).isoformat())
        self.assertEqual([r["item"] for r in res["rows"]], [self.vaccine.id])

    def test_a_file_that_is_not_a_spreadsheet_is_refused(self):
        with self.assertRaises(PriceRowError):
            parse_price_upload(SimpleUploadedFile("prices.pdf", b"%PDF"), default_date=None)

    def test_the_template_lists_every_item_with_its_current_price(self):
        self.price(self.starter, 42, -5)
        response = self.client.get(reverse("item_price_list_template_data"))
        self.assertEqual(response.status_code, 200)
        rows = list(openpyxl.load_workbook(io.BytesIO(response.content)).active
                    .iter_rows(values_only=True))
        self.assertIn("New Price", rows[0])
        starter = next(r for r in rows if r[0] == self.starter.item_code)
        self.assertEqual(starter[4], 42)
        self.assertEqual(len(rows) - 1, Item.objects.count())

    def test_the_page_previews_an_upload(self):
        response = self.client.post(reverse("item_price_list_upload_data"), {
            "file": self.workbook([(self.vaccine.item_code, 5, None)]),
            "effective_date": self.day(1).isoformat(),
        })
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()["counts"]["create"], 1)
        self.assertFalse(ItemPriceList.objects.exists())


class AccessMappingTests(TestCase):

    def test_every_new_endpoint_is_governed_by_the_price_list_tab(self):
        """Named so the existing rule maps them: reading and previewing need
        View, saving needs Add. An unmapped url would be refused to everyone
        once web access is enforced."""
        from user.access import derive_tab, resolve_action

        expected = {
            "item_price_list_overview_data": "view",
            "item_price_list_history_data": "view",
            "item_price_list_audit_data": "view",
            "item_price_list_revise_data": "view",
            "item_price_list_upload_data": "view",
            "item_price_list_template_data": "view",
            "item_price_list_bulk_create": "add",
        }
        for name, action in expected.items():
            self.assertEqual(resolve_action(name) or derive_tab(name),
                             ("item_price_list", action), name)


class LastPurchaseRateTests(PriceListBase):
    """The rate on the item's most recent purchase bill, beside its price."""

    def setUp(self):
        super().setUp()
        from account.models import AccountType, ChartOfAccount, CompanyProfile
        from inventory.models import Warehouse
        from purchase.models import Supplier

        self.store = Warehouse.objects.create(name="Bahraich Warehouse")
        self.mills = Supplier.objects.create(name="Maharashtra Feeds Pvt Ltd")
        self.breeder = Supplier.objects.create(name="Ganga Breeding Farm")
        account_type = AccountType.objects.create(
            name="T", code_range_start=500000, code_range_end=599999, report="PL")
        self.pay_account = ChartOfAccount.objects.create(
            company=CompanyProfile.get_solo(), code="500001", description="X",
            account_type=account_type)

    def buy(self, item, rate, days, unit="Bag"):
        from purchase.models import GeneralPurchase, GeneralPurchaseItem

        purchase = GeneralPurchase.objects.create(date=self.day(days), supplier=self.mills)
        GeneralPurchaseItem.objects.create(
            purchase=purchase, item=item, farm_warehouse=self.store, unit=unit,
            rcv_qty=Decimal("100"), rate=Decimal(str(rate)),
            discount_percent=Decimal("0"), discount_amount=Decimal("0"),
            gst_percent=Decimal("0"))
        return purchase

    def buy_chicks(self, item, rate, days):
        from purchase.models import ChicksPurchase, ChicksPurchaseItem

        purchase = ChicksPurchase.objects.create(date=self.day(days), supplier=self.breeder, item=item)
        ChicksPurchaseItem.objects.create(purchase=purchase, farm_warehouse=self.store,
                                          sent_qty=Decimal("1000"), rate=Decimal(str(rate)))
        return purchase

    def buy_eggs(self, item, rate, days):
        from hatchery.models import EggPurchase, EggPurchaseItem

        purchase = EggPurchase.objects.create(date=self.day(days), supplier=self.breeder,
                                              warehouse=self.store, pay_account=self.pay_account)
        EggPurchaseItem.objects.create(egg_purchase=purchase, item=item,
                                       rcv_qty=Decimal("500"), rate=Decimal(str(rate)))
        return purchase

    def last(self, item):
        from inventory.services.price_list import last_purchase_rates
        return last_purchase_rates([item.id]).get(item.id)

    def test_the_latest_bill_wins_with_its_own_unit(self):
        self.buy(self.starter, 40, -30)
        latest = self.buy(self.starter, 44, -3, unit="Kg")
        self.assertEqual(self.last(self.starter), {
            "rate": "44.00", "unit": "Kg", "date": self.day(-3).isoformat(),
            "supplier": "Maharashtra Feeds Pvt Ltd", "ref": latest.purchase_no,
            "source": "General Purchase",
        })

    def test_on_the_same_day_the_bill_entered_last_wins(self):
        self.buy(self.starter, 40, -2)
        self.buy(self.starter, 41, -2)
        self.assertEqual(self.last(self.starter)["rate"], "41.00")

    def test_a_bill_with_no_rate_is_passed_over(self):
        """A free or sample delivery is not a price of nothing."""
        self.buy(self.starter, 40, -10)
        self.buy(self.starter, 0, -1)
        self.assertEqual(self.last(self.starter)["rate"], "40.00")

    def test_chicks_and_egg_purchases_count_too(self):
        self.buy(self.vaccine, 5, -10)
        self.buy_chicks(self.vaccine, 6, -5)
        self.buy_eggs(self.tonic, "4.50", -2)
        self.assertEqual((self.last(self.vaccine)["rate"], self.last(self.vaccine)["source"]),
                         ("6.00", "Chicks Purchase"))
        self.assertEqual((self.last(self.tonic)["rate"], self.last(self.tonic)["source"]),
                         ("4.50", "Egg Purchase"))

    def test_the_list_and_the_history_carry_it(self):
        self.buy(self.starter, 44, -3)
        rows = price_overview(today=self.today)
        self.assertEqual(self.row_for(rows, self.starter)["last_purchase"]["rate"], "44.00")
        self.assertIsNone(self.row_for(rows, self.tonic)["last_purchase"])
        history = item_price_history(self.starter, today=self.today)
        self.assertEqual(history["last_purchase"]["rate"], "44.00")
