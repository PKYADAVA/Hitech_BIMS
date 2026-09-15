"""Items can be made Active or Inactive.

An inactive item drops out of the item pickers on new entries, and nothing
already recorded against it changes. The one place that needs care is an edit
screen: a record that already uses an item made inactive since must still be
offered that item, or its own row comes up blank and saving it loses the item.
"""
import io
import json
from datetime import timedelta
from decimal import Decimal

import openpyxl
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from inventory.models import Item, ItemCategory, ItemPriceList, Warehouse
from inventory.services.price_list import (parse_price_upload,
                                           price_overview,
                                           price_template_workbook,
                                           revise_preview)


class ItemActiveBase(TestCase):

    def setUp(self):
        self.today = timezone.localdate()
        feed = ItemCategory.objects.create(name="Broiler Feed")
        common = dict(valuation_method="Weighted Average", usage="Produced",
                      source="Purchased", type="Raw Material", item_account="Expense",
                      standard_cost_per_unit=0)
        self.starter = Item.objects.create(description="Starter Feed", category=feed, **common)
        self.finisher = Item.objects.create(description="Finisher Feed", category=feed, **common)
        self.user = get_user_model().objects.create_superuser("keeper", "k@x.com", "Str0ngPass!")
        self.client.force_login(self.user)

    def retire(self, item):
        item.is_active = False
        item.save(update_fields=["is_active"])


class ToggleTests(ItemActiveBase):

    def test_a_new_item_starts_active(self):
        self.assertTrue(self.starter.is_active)

    def test_the_button_flips_the_status_both_ways(self):
        url = reverse("items_toggle_active", args=[self.finisher.id])
        first = self.client.post(url).json()
        self.assertFalse(first["is_active"])
        self.assertIn("Inactive", first["message"])
        self.finisher.refresh_from_db()
        self.assertFalse(self.finisher.is_active)
        self.assertTrue(self.client.post(url).json()["is_active"])

    def test_it_only_changes_on_a_post(self):
        response = self.client.get(reverse("items_toggle_active", args=[self.finisher.id]))
        self.assertEqual(response.status_code, 405)
        self.finisher.refresh_from_db()
        self.assertTrue(self.finisher.is_active)

    def test_it_needs_edit_rights_on_items(self):
        from user.access import derive_tab, resolve_action
        name = "items_toggle_active"
        self.assertEqual(resolve_action(name) or derive_tab(name), ("items", "edit"))

    def test_the_items_list_reports_the_status(self):
        self.retire(self.finisher)
        rows = {r["id"]: r["is_active"] for r in self.client.get(reverse("item_list")).json()}
        self.assertEqual(rows, {self.starter.id: True, self.finisher.id: False})


class PickerTests(ItemActiveBase):

    def test_new_entries_are_offered_active_items_only(self):
        self.retire(self.finisher)
        self.assertEqual(list(Item.objects.for_entry()), [self.starter])

    def test_a_record_being_edited_keeps_its_own_inactive_item(self):
        self.retire(self.finisher)
        offered = set(Item.objects.for_entry(keep=[self.finisher.id]))
        self.assertEqual(offered, {self.starter, self.finisher})

    def test_the_daily_entry_form_hides_an_inactive_feed(self):
        self.retire(self.finisher)
        items = list(self.client.get(reverse("daily_entry_add")).context["items"])
        self.assertIn(self.starter, items)
        self.assertNotIn(self.finisher, items)

    def test_the_medicine_entry_form_hides_an_inactive_item(self):
        self.retire(self.finisher)
        items = list(self.client.get(reverse("medicine_entry_add")).context["items"])
        self.assertNotIn(self.finisher, items)

    def test_a_purchase_being_edited_still_offers_its_inactive_item(self):
        from purchase.models import GeneralPurchase, GeneralPurchaseItem, Supplier
        from purchase.views import _general_purchase_form_context

        store = Warehouse.objects.create(name="Bahraich Warehouse")
        purchase = GeneralPurchase.objects.create(
            date=self.today - timedelta(days=5),
            supplier=Supplier.objects.create(name="Maharashtra Feeds Pvt Ltd"))
        GeneralPurchaseItem.objects.create(
            purchase=purchase, item=self.finisher, farm_warehouse=store,
            rcv_qty=Decimal("100"), rate=Decimal("42"), discount_percent=Decimal("0"),
            discount_amount=Decimal("0"), gst_percent=Decimal("0"))
        self.retire(self.finisher)

        editing = set(_general_purchase_form_context(self.user, purchase)["items"])
        adding = set(_general_purchase_form_context(self.user, None)["items"])
        self.assertIn(self.finisher, editing)
        self.assertNotIn(self.finisher, adding)

    def test_reports_still_list_every_item(self):
        """History does not change when an item goes out of use."""
        self.retire(self.finisher)
        items = list(self.client.get(reverse("stock_transfer_list")).context["items"])
        self.assertIn(self.finisher, items)


class PriceListTests(ItemActiveBase):

    def setUp(self):
        super().setUp()
        for item in (self.starter, self.finisher):
            ItemPriceList.objects.create(item=item, price=Decimal("42"),
                                         effective_date=self.today - timedelta(days=5))
        self.retire(self.finisher)

    def test_it_reads_inactive_and_keeps_its_price(self):
        """One status per row, so an inactive item never reads Active too."""
        row = next(r for r in price_overview(today=self.today) if r["item"] == self.finisher.id)
        self.assertEqual((row["status"], row["status_label"]), ("inactive", "Inactive"))
        self.assertEqual(row["price"], "42.00")
        self.assertFalse(row["is_active"])

    def test_the_full_list_shows_active_and_inactive_together(self):
        res = self.client.get(reverse("item_price_list_overview_data")).json()
        self.assertEqual({r["item"]: r["status"] for r in res["rows"]},
                         {self.starter.id: "active", self.finisher.id: "inactive"})
        self.assertEqual((res["counts"]["active"], res["counts"]["inactive"]), (1, 1))

    def test_add_prices_offers_active_items_only(self):
        response = self.client.get(reverse("item_price_list"))
        ids = [i["id"] for i in response.context["items_json"]]
        self.assertEqual(ids, [self.starter.id])
        self.assertTrue(response.context["can_toggle_items"])

    def test_bulk_revision_skips_it(self):
        rows = revise_preview([self.starter.id, self.finisher.id], "percent", 10,
                              self.today + timedelta(days=1))["rows"]
        by_item = {r["item"]: r for r in rows}
        self.assertEqual(by_item[self.starter.id]["action"], "create")
        self.assertEqual(by_item[self.finisher.id]["action"], "skip")
        self.assertEqual(by_item[self.finisher.id]["message"], "Item is inactive")

    def test_an_upload_flags_it(self):
        book = openpyxl.Workbook()
        book.active.append(["Item Code", "New Price"])
        book.active.append([self.finisher.item_code, 45])
        buffer = io.BytesIO()
        book.save(buffer)
        row = parse_price_upload(SimpleUploadedFile("prices.xlsx", buffer.getvalue()),
                                 default_date=self.today.isoformat())["rows"][0]
        self.assertEqual((row["action"], row["message"]), ("error", "Item is inactive"))

    def test_the_template_leaves_it_out(self):
        sheet = openpyxl.load_workbook(io.BytesIO(price_template_workbook(today=self.today))).active
        codes = [r[0] for r in sheet.iter_rows(min_row=2, values_only=True)]
        self.assertEqual(codes, [self.starter.item_code])
